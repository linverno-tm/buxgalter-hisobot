# -*- coding: utf-8 -*-
"""
bh_report - "КАМЕРАЛ ТЕКШИРУВЛАР" ko'rinishidagi Excel hisobotini yozadi.

Varaqlar mavjud fayldagidek:
    касса YYYY  - kassa cheklari (20 ustun)
    YYYY ТХ     - tovar harakati (22 ustun, ikki qavatli sarlavha)
    Xatolar     - YANGI: tekshiruv natijalari
    Ma'lumot    - YANGI: hisobot qanday yig'ilgani

Chiqish formati .xlsx. Eski fayl .xls (BIFF8) edi, lekin Python'da .xls
yozish kutubxonalari tashlab yuborilgan. .xlsx Excel'da bir xil ochiladi.
"""

import os
import datetime
from decimal import Decimal

import bh_config as C
import bh_db as DB
import bh_fifo as F


NUM_MONEY = '#,##0.00'
NUM_QTY = '#,##0.####'
NUM_DATE = 'DD.MM.YYYY'


class StyleBook:
    """
    NamedStyle keshi.

    openpyxl'da har `cell.font = ...` topshirig'i stilni xeshlab ish
    kitobining stil jadvalidan qidiradi. 10 000 satrli kassa varag'ida bu
    150 000 dan ortiq qidiruv - hisobotning eng sekin joyi aynan shu edi.

    NamedStyle bir marta ro'yxatdan o'tkaziladi, keyin `cell.style = nom`
    faqat tayyor indeksni ko'chiradi: ko'rinish bir xil, tezlik ~5 barobar.

    name(asos, raqam_formati, rang) -> stil nomi. Rang ("warn" / "err")
    asosning fonini almashtiradi, shuning uchun satrni bo'yash uchun
    ikkinchi marta aylanib chiqish shart emas.
    """

    def __init__(self, wb):
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        self.wb = wb
        self._names = {}

        thin = Side(style="thin", color="9AA0A6")
        med = Side(style="medium", color="44546A")
        box = Border(left=thin, right=thin, top=thin, bottom=thin)

        self.font_cell = Font(name="Calibri", size=9)
        self.font_total = Font(name="Calibri", size=10, bold=True)

        self.tints = {
            "warn": PatternFill("solid", fgColor="FFF2CC"),
            "err": PatternFill("solid", fgColor="FCE4E4"),
        }

        self.bases = {
            "title": {"font": Font(name="Calibri", size=14, bold=True, color="1F3864"),
                      "align": Alignment(horizontal="center", vertical="center")},
            "hdr": {"font": Font(name="Calibri", size=9, bold=True, color="FFFFFF"),
                    "fill": PatternFill("solid", fgColor="44546A"),
                    "align": Alignment(horizontal="center", vertical="center",
                                       wrap_text=True),
                    "border": Border(left=thin, right=thin, top=med, bottom=thin)},
            "sub": {"font": Font(name="Calibri", size=9, bold=True, color="FFFFFF"),
                    "fill": PatternFill("solid", fgColor="5B7397"),
                    "align": Alignment(horizontal="center", vertical="center",
                                       wrap_text=True),
                    "border": Border(left=thin, right=thin, top=thin, bottom=med)},
            "cell": {"font": self.font_cell,
                     "align": Alignment(vertical="top", wrap_text=True),
                     "border": box},
            "num": {"font": self.font_cell,
                    "align": Alignment(horizontal="right", vertical="top"),
                    "border": box},
            "total": {"font": self.font_total,
                      "fill": PatternFill("solid", fgColor="D9E2F3"),
                      "align": Alignment(horizontal="right", vertical="center"),
                      "border": Border(left=thin, right=thin, top=med, bottom=med)},
        }

    def name(self, base, numfmt=None, tint=None):
        key = (base, numfmt or "", tint or "")
        nm = self._names.get(key)
        if nm is not None:
            return nm

        from openpyxl.styles import NamedStyle

        spec = self.bases[base]
        nm = "bh%d" % len(self._names)
        ns = NamedStyle(name=nm)
        if "font" in spec:
            ns.font = spec["font"]
        fill = self.tints[tint] if tint else spec.get("fill")
        if fill is not None:
            ns.fill = fill
        if "align" in spec:
            ns.alignment = spec["align"]
        if "border" in spec:
            ns.border = spec["border"]
        if numfmt:
            ns.number_format = numfmt
        self.wb.add_named_style(ns)
        self._names[key] = nm
        return nm


def _f(v):
    """Decimal -> Excel raqami (float). Excel Decimal'ni tushunmaydi."""
    if v is None:
        return None
    if isinstance(v, Decimal):
        return float(v)
    return v


# ===========================================================================
# ТХ varag'i
# ===========================================================================
def write_tx_sheet(wb, cx, year, owner_name, S):
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("%d ТХ" % year)
    rows = F.year_rows(cx, year)
    tot = F.year_totals(rows)
    ncol = len(C.TX_COLUMNS)

    # --- 1-satr: sarlavha ---
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=ncol)
    c = ws.cell(row=1, column=2, value="%s ning %d йил товар хисоботи" % (owner_name, year))
    c.style = S.name("title")
    ws.row_dimensions[1].height = 26

    # --- 2-satr: jami ko'rsatkichlar ---
    ws.cell(row=2, column=11, value="Кирим жами:")
    ws.cell(row=2, column=12, value=_f(tot["in_sum"])).number_format = NUM_MONEY
    ws.cell(row=2, column=14, value="Чиқим жами:")
    ws.cell(row=2, column=15, value=_f(tot["out_sum"])).number_format = NUM_MONEY
    ws.cell(row=2, column=17, value="Қолдиқ:")
    ws.cell(row=2, column=18, value=_f(tot["close_sum"])).number_format = NUM_MONEY
    ws.cell(row=2, column=20, value="Соф фойда:")
    ws.cell(row=2, column=21, value=_f(tot["profit"])).number_format = NUM_MONEY
    for col in (11, 14, 17, 20):
        ws.cell(row=2, column=col).font = S.font_total

    # --- 3-4 satr: ikki qavatli sarlavha ---
    hr, sr = 3, 4
    for i, (_idx, name, width) in enumerate(C.TX_COLUMNS):
        col = i + 1
        ws.column_dimensions[get_column_letter(col)].width = width
    # guruhlanmagan ustunlar 3-4 satrni egallaydi
    grouped = set()
    for a, b, _t in C.TX_GROUPS:
        grouped.update(range(a, b + 1))
    for i, (idx, name, _w) in enumerate(C.TX_COLUMNS):
        col = i + 1
        if idx in grouped:
            continue
        ws.merge_cells(start_row=hr, start_column=col, end_row=sr, end_column=col)
        ws.cell(row=hr, column=col, value=name).style = S.name("hdr")
    for a, b, title in C.TX_GROUPS:
        ws.merge_cells(start_row=hr, start_column=a + 1, end_row=hr, end_column=b + 1)
        ws.cell(row=hr, column=a + 1, value=title).style = S.name("hdr")
        for idx in range(a, b + 1):
            ws.cell(row=sr, column=idx + 1,
                    value=C.TX_COLUMNS[idx][1]).style = S.name("sub")
    ws.row_dimensions[hr].height = 34
    ws.row_dimensions[sr].height = 18

    # --- ma'lumot satrlari ---
    # Har ustunning stili BIR MARTA hisoblanadi, rangli satr uchun tayyor
    # variant olinadi - satrni bo'yash uchun ikkinchi marta aylanish yo'q.
    def _row_styles(tint):
        out = []
        for i in range(ncol):
            if i == 1:
                out.append(S.name("cell", NUM_DATE, tint))
            elif i == 0:
                out.append(S.name("num", "0", tint))
            elif i in (7, 10, 13, 16, 19):
                out.append(S.name("num", NUM_QTY, tint))
            elif i >= 6:
                out.append(S.name("num", NUM_MONEY, tint))
            else:
                out.append(S.name("cell", None, tint))
        return out

    st_plain, st_warn, st_err = (_row_styles(None), _row_styles("warn"),
                                 _row_styles("err"))
    r = sr + 1
    for n, d in enumerate(rows, start=1):
        vals = [
            n,
            C.parse_date(d["date"]),
            d["name"],
            ("маркировкаланган" if d["is_marked"] else
             ("маркировкасиз" if d["is_marked"] is False else "")),
            ("%s - %s" % (d["mxik"], d["mxik_name"])).strip(" -") if d["mxik"] else "",
            d["unit"],
            _f(d["sale_price"]),
            _f(d["open_qty"]), _f(d["cost"]), _f(d["open_sum"]),
            _f(d["in_qty"]), _f(d["cost"]), _f(d["in_sum"]),
            _f(d["out_qty"]), _f(d["cost"]), _f(d["out_sum"]),
            _f(d["close_qty"]), _f(d["cost"]), _f(d["close_sum"]),
            _f(d["sale_qty"]),
            _f(d["sale_price"]), _f(d["sale_sum"]),
        ]
        if d.get("shortfall"):
            sty = st_err
        elif d["close_qty"] < 0:
            sty = st_warn
        else:
            sty = st_plain
        for i, v in enumerate(vals):
            ws.cell(row=r, column=i + 1, value=v).style = sty[i]
        r += 1

    # --- jami satri ---
    tot_plain = S.name("total")
    tot_qty = S.name("total", NUM_QTY)
    tot_money = S.name("total", NUM_MONEY)
    ws.cell(row=r, column=1, value="ЖАМИ").style = tot_plain
    for col in range(2, ncol + 1):
        ws.cell(row=r, column=col).style = tot_plain
    pairs = [(8, tot["open_qty"]), (10, tot["open_sum"]),
             (11, tot["in_qty"]), (13, tot["in_sum"]),
             (14, tot["out_qty"]), (16, tot["out_sum"]),
             (17, tot["close_qty"]), (19, tot["close_sum"]),
             (20, tot["out_qty"]), (22, tot["sale_sum"])]
    for col, val in pairs:
        cell = ws.cell(row=r, column=col, value=_f(val))
        cell.style = tot_qty if col in (8, 11, 14, 17, 20) else tot_money

    ws.freeze_panes = ws.cell(row=sr + 1, column=4)
    if r > sr + 1:
        ws.auto_filter.ref = "A%d:%s%d" % (sr, get_column_letter(ncol), r - 1)
    ws.sheet_view.zoomScale = 90
    return len(rows), tot


# ===========================================================================
# Kassa varag'i
# ===========================================================================
def write_kassa_sheet(wb, cx, year, S, progress=None):
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("касса %d" % year)
    widths = [34, 16, 16, 19, 11, 46, 10, 15, 13, 13, 13, 15, 15, 14, 19, 13,
              14, 16, 20, 11]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w

    for i, h in enumerate(C.KASSA_HEADERS):
        ws.cell(row=1, column=i + 1, value=h).style = S.name("hdr")
    ws.row_dimensions[1].height = 32

    # Ustun stillari oldindan (qaytarish satrlari uchun sariq variant ham)
    def _kassa_styles(tint):
        out = []
        for i in range(len(C.KASSA_HEADERS)):
            if i == 6:
                out.append(S.name("num", NUM_QTY, tint))
            elif i in (7, 8, 9, 10, 11, 12, 13):
                out.append(S.name("num", NUM_MONEY, tint))
            else:
                out.append(S.name("cell", None, tint))
        return out

    st_plain, st_warn = _kassa_styles(None), _kassa_styles("warn")
    r = 2
    n = 0
    for d in cx.execute("""
            SELECT l.*, d.doc_no, d.doc_date, d.pos_id, d.partner_tin, d.check_type,
                   d.total_gross, d.total_vat, d.is_return
            FROM doc_line l JOIN document d ON d.id=l.document_id
            WHERE l.kind='chiqim' AND d.doc_year=?
            ORDER BY d.doc_date, d.doc_no, l.line_no""", (year,)):
        cash = DB.D(0)
        card = DB.D(d["total_gross"])
        vals = [
            "", d["partner_tin"], d["pos_id"], d["doc_date"], d["doc_no"],
            d["raw_name"], _f(DB.D(d["qty"])), _f(DB.D(d["amount_gross"])),
            0, 0, _f(DB.D(d["vat_amount"])),
            _f(cash), _f(card), _f(DB.D(d["total_vat"])),
            d["mxik"], d["unit_raw"], d["barcode"], "", d["marking_code"],
            d["check_type"],
        ]
        sty = st_warn if d["is_return"] else st_plain
        for i, v in enumerate(vals):
            ws.cell(row=r, column=i + 1, value=v).style = sty[i]
        r += 1
        n += 1
        if progress and n % 2000 == 0:
            progress(n)

    ws.freeze_panes = "A2"
    if n:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(C.KASSA_HEADERS)), r - 1)
    ws.sheet_view.zoomScale = 90
    return n


# ===========================================================================
# Xatolar varag'i (YANGI)
# ===========================================================================
def write_issues_sheet(wb, cx, years, S):
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("Xatolar")
    heads = ["Yil", "Daraja", "Turi", "Tavsif", "Manba", "ID"]
    widths = [8, 16, 26, 96, 12, 8]
    for i, (h, w) in enumerate(zip(heads, widths)):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.cell(row=1, column=i + 1, value=h).style = S.name("hdr")
    ws.row_dimensions[1].height = 22

    r = 2
    q = "SELECT * FROM issue WHERE resolved=0"
    args = []
    if years:
        q += " AND (year IS NULL OR year IN (%s))" % ",".join("?" * len(years))
        args = list(years)
    q += (" ORDER BY CASE severity WHEN 'xato' THEN 0 WHEN 'ogohlantirish' "
          "THEN 1 ELSE 2 END, year, code")
    n = 0
    iss_plain = S.name("cell")
    iss_warn = S.name("cell", None, "warn")
    iss_err = S.name("cell", None, "err")
    for d in cx.execute(q, args):
        vals = [d["year"], d["severity"],
                C.ISSUE_TITLES.get(d["code"], d["code"]), d["message"],
                d["ref_table"], d["ref_id"]]
        if d["severity"] == C.SEVERITY_ERROR:
            sty = iss_err
        elif d["severity"] == C.SEVERITY_WARN:
            sty = iss_warn
        else:
            sty = iss_plain
        for i, v in enumerate(vals):
            ws.cell(row=r, column=i + 1, value=v).style = sty
        r += 1
        n += 1

    ws.freeze_panes = "A2"
    if n:
        ws.auto_filter.ref = "A1:F%d" % (r - 1)
    return n


# ===========================================================================
# Ma'lumot varag'i (YANGI)
# ===========================================================================
def write_info_sheet(wb, cx, years, S, totals_by_year):
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet("Ma'lumot", 0)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 22

    r = 1
    ws.cell(row=r, column=1, value="Hisobot qanday yig'ilgani").style = S.name("title")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 2

    st = DB.stats(cx)
    info = [
        ("Yaratilgan sana", datetime.datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Dastur versiyasi", C.VERSION),
        ("QQS stavkasi", "%s%%" % (C.VAT_RATE * 100)),
        ("", ""),
        ("Import qilingan fayl", st["files"]),
        ("Kirim hujjati (faktura)", st["docs_in"]),
        ("Chiqim hujjati (chek)", st["docs_out"]),
        ("Kirim satri", st["lines_in"]),
        ("Chiqim satri", st["lines_out"]),
        ("Mahsulot kartochkasi", st["products"]),
        ("Bog'lanmagan satr", st["unmatched"]),
        ("Ochiq tekshiruv yozuvi", st["issues"]),
    ]
    for k, v in info:
        if k:
            ws.cell(row=r, column=1, value=k).font = S.font_cell
            ws.cell(row=r, column=2, value=v).font = S.font_total
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Yillar bo'yicha").style = S.name("hdr")
    for i, h in enumerate(["Kirim (tannarx)", "Chiqim (tannarx)",
                           "Qoldiq (tannarx)", "Sotuv (QQS bilan)"]):
        ws.cell(row=r, column=i + 2, value=h).style = S.name("hdr")
    r += 1
    for y in years:
        t = totals_by_year.get(y)
        if not t:
            continue
        ws.cell(row=r, column=1, value=y).font = S.font_total
        for i, k in enumerate(("in_sum", "out_sum", "close_sum", "sale_sum")):
            c = ws.cell(row=r, column=i + 2, value=_f(t[k]))
            c.number_format = NUM_MONEY
        r += 1

    r += 1
    notes = [
        "Eslatmalar:",
        "  - Qizil bo'yalgan satr: tovar sotilgan, lekin kirim hujjati topilmadi "
        "(tannarx taxminiy).",
        "  - Sariq bo'yalgan satr: davr oxiriga qoldiq manfiy yoki chek qaytarilgan.",
        "  - Davr boshiga qoldiq o'tgan yildan avtomatik ko'chiriladi (FIFO).",
        "  - Chiqim tannarxi FIFO bo'yicha: eng eski partiyadan yechiladi.",
        "  - Kassa 'Нархи' ustuni satr SUMMASI (QQS ichida), birlik narxi emas.",
    ]
    for t in notes:
        ws.cell(row=r, column=1, value=t).font = S.font_cell
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        r += 1
    return ws


# ===========================================================================
# Asosiy
# ===========================================================================
def generate(cx, out_path, years=None, owner_name=None, progress=None):
    """
    Hisobotni yozadi. (yo'l, statistika) qaytaradi.
    """
    import openpyxl

    years = years or DB.available_years(cx)
    years = sorted(y for y in years if y)
    if not years:
        raise ValueError("Hisobot uchun ma'lumot yo'q - avval fayllarni import qiling.")

    owner_name = owner_name or DB.get_setting(cx, "owner_name", "Ташкилот")

    # Hisobotdan oldin ombor majburan qayta hisoblanadi: qolda
    # boglangan sotuvlar ham qoldiqdan ayrilsin (Qayta hisoblash
    # tugmasi bosilmagan bolsa ham).
    import bh_matching as _M
    if progress:
        progress(0, 0, 'ombor yangilanmoqda')
    try:
        _eng = _M.MatchEngine(cx)
        _M.auto_match_all(cx, _eng)
    except Exception:
        pass
    F.rebuild_stock(cx)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    S = StyleBook(wb)

    stats = {"years": {}, "kassa_rows": 0, "tx_rows": 0}
    totals_by_year = {}
    steps = len(years) * 2 + 2
    step = 0

    for y in years:
        F.validate_year(cx, y)
        n = write_kassa_sheet(
            wb, cx, y, S,
            progress=(lambda k, _y=y, _s=step: progress(_s, steps, "касса %d (%d satr)" % (_y, k)))
            if progress else None)
        stats["kassa_rows"] += n
        step += 1
        if progress:
            progress(step, steps, "касса %d" % y)

    for y in years:
        nrows, tot = write_tx_sheet(wb, cx, y, owner_name, S)
        stats["tx_rows"] += nrows
        stats["years"][y] = tot
        totals_by_year[y] = tot
        step += 1
        if progress:
            progress(step, steps, "%d ТХ" % y)

    n_iss = write_issues_sheet(wb, cx, years, S)
    stats["issues"] = n_iss
    step += 1
    if progress:
        progress(step, steps, "Xatolar")

    write_info_sheet(wb, cx, years, S, totals_by_year)
    step += 1
    if progress:
        progress(step, steps, "Ma'lumot")

    d = os.path.dirname(os.path.abspath(out_path))
    if d and not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)
    wb.save(out_path)
    wb.close()
    return out_path, stats


def default_filename(years):
    ys = sorted(y for y in (years or []) if y)
    if not ys:
        rng = "hisobot"
    elif len(ys) == 1:
        rng = "%d" % ys[0]
    else:
        rng = "%d-%d" % (ys[0], ys[-1])
    return "КАМЕРАЛ ТЕКШИРУВЛАР %s.xlsx" % rng
